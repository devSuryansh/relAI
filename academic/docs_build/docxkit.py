"""Write institute-format Word documents on top of an existing template.

The template keeps its page size, margins, borders, headers, footers, and the
page-number format of each section. The body is replaced. Contents pages are
filled in two passes: render to PDF, read where each heading landed, rebuild.
"""

from __future__ import annotations

import copy
import re
import subprocess
import tempfile
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

FONT = "Times New Roman"
MONO = "Courier New"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


@dataclass
class Entry:
    """A heading or caption that a contents page points at."""

    kind: str
    level: int
    text: str
    label: str
    zone: str


@dataclass
class Style:
    body: float = 12
    line: float = 1.15
    after: float = 6
    h1: float = 14
    h2: float = 12
    h3: float = 12
    caption: float = 11
    table: float = 11
    h1_upper: bool = True
    h1_color: RGBColor = field(default_factory=lambda: BLACK)
    numbered_heading: bool = False


class Builder:
    def __init__(self, template: Path, style: Style, pages: dict[str, str] | None = None) -> None:
        self.doc = Document(str(template))
        self.style = style
        self.pages = pages or {}
        self.entries: list[Entry] = []
        self.zone = "front"
        self.chapter_no = 0
        self.section_no = 0
        self.figure_no = 0
        self.table_no = 0
        self.logo = self._extract_logo()
        self.sect_prs = [copy.deepcopy(section._sectPr) for section in self.doc.sections]
        body = self.doc.element.body
        final = body[-1] if body[-1].tag == qn("w:sectPr") else None
        for child in list(body):
            if child is not final:
                body.remove(child)
        self._fresh_page = True

    # ----- low level -------------------------------------------------------------

    def paragraph(
        self,
        text: str = "",
        *,
        size: float | None = None,
        bold: bool = False,
        italic: bool = False,
        align: str = "justify",
        before: float = 0,
        after: float | None = None,
        line: float | None = None,
        color: RGBColor | None = None,
        indent: float | None = None,
        hanging: float | None = None,
        font: str = FONT,
        keep_next: bool = False,
    ):
        p = self.doc.add_paragraph()
        fmt = p.paragraph_format
        p.alignment = ALIGN[align]
        fmt.space_before = Pt(before)
        fmt.space_after = Pt(self.style.after if after is None else after)
        fmt.line_spacing = self.style.line if line is None else line
        if indent is not None:
            fmt.left_indent = Inches(indent)
        if hanging is not None:
            fmt.first_line_indent = Inches(-hanging)
            fmt.tab_stops.add_tab_stop(Inches(indent or hanging))
        fmt.keep_with_next = keep_next
        self.runs(p, text, size=size or self.style.body, bold=bold, italic=italic, color=color, font=font)
        self._fresh_page = False
        return p

    def runs(self, p, text: str, *, size, bold=False, italic=False, color=None, font=FONT) -> None:
        """Add text with **bold**, *italic*, and `code` spans."""
        for token in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
            if not token:
                continue
            b, i, f, t = bold, italic, font, token
            if token.startswith("**") and token.endswith("**"):
                b, t = True, token[2:-2]
            elif token.startswith("*") and token.endswith("*") and len(token) > 2:
                i, t = True, token[1:-1]
            elif token.startswith("`") and token.endswith("`"):
                f, t = MONO, token[1:-1]
            run = p.add_run(t)
            set_font(run, f, size if f == font else size - 1, b, i, color)

    def page_break(self) -> None:
        if self._fresh_page:
            return
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_break(WD_BREAK.PAGE)
        self._fresh_page = True

    def end_section(self, index: int) -> None:
        """Close the current section with the template's section properties."""
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p._p.get_or_add_pPr().append(copy.deepcopy(self.sect_prs[index]))
        self._fresh_page = True

    # ----- structure ---------------------------------------------------------------

    def front_heading(self, text: str, *, size: float | None = None, record: bool = True) -> None:
        self.page_break()
        self.paragraph(text, size=size or self.style.h1, bold=True, align="center", line=1.5, after=12)
        if record:
            self.entries.append(Entry("front", 1, text.title(), text, "front"))

    def chapter(self, title: str) -> None:
        self.page_break()
        self.zone = "body"
        self.chapter_no += 1
        self.section_no = 0
        self.figure_no = 0
        self.table_no = 0
        shown = title.upper() if self.style.h1_upper else title
        if self.style.numbered_heading:
            shown = f"{self.chapter_no}. {shown}"
            self.paragraph(shown, size=self.style.h1, bold=True, align="center", line=1.3, after=12, color=self.style.h1_color, keep_next=True)
            self.entries.append(Entry("chapter", 1, shown, shown, "body"))
            return
        label = f"CHAPTER {self.chapter_no}"
        self.paragraph(label, size=self.style.h1, bold=True, align="center", line=1.2, after=2, color=self.style.h1_color, keep_next=True)
        self.paragraph(shown, size=self.style.h1, bold=True, align="center", line=1.3, after=12, color=self.style.h1_color, keep_next=True)
        self.entries.append(Entry("chapter", 1, f"{label}: {title.upper()}", shown, "body"))

    def unnumbered(self, title: str, *, new_page: bool = True) -> None:
        if new_page:
            self.page_break()
        self.paragraph(title.upper(), size=self.style.h1, bold=True, align="center", line=1.3, after=12, color=self.style.h1_color, keep_next=True)
        self.entries.append(Entry("chapter", 1, title.upper(), title.upper(), "body"))

    def section(self, title: str) -> None:
        self.section_no += 1
        text = f"{self.chapter_no}.{self.section_no} {title}"
        self.paragraph(text, size=self.style.h2, bold=True, align="left", before=8, after=4, keep_next=True)
        self.entries.append(Entry("section", 2, text, text, "body"))

    def subsection(self, title: str) -> None:
        self.paragraph(title, size=self.style.h3, bold=True, italic=False, align="left", before=4, after=3, keep_next=True)

    def bullets(self, items: Sequence[str], *, marker: str = "\u2022", indent: float = 0.35) -> None:
        for item in items:
            self.paragraph(f"{marker}\t{item}", indent=indent, hanging=0.25, after=3)

    def numbered(self, items: Sequence[str], *, indent: float = 0.4) -> None:
        for n, item in enumerate(items, 1):
            self.paragraph(f"{n}.\t{item}", indent=indent, hanging=0.3, after=3)

    def figure(self, path: Path, caption: str, *, width: float = 6.0, max_height: float = 7.6) -> None:
        w, h = Image.open(path).size
        width = min(width, max_height * w / h)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Inches(width))
        self.figure_no += 1
        label = f"Figure {self.chapter_no}.{self.figure_no}: {caption}"
        self.paragraph(label, size=self.style.caption, italic=True, align="center", after=10, line=1.0)
        self.entries.append(Entry("figure", 1, label, label, "body"))

    def table(
        self,
        rows: Sequence[Sequence[str]],
        widths: Sequence[float],
        caption: str | None = None,
        *,
        size: float | None = None,
        header: bool = True,
        align: Sequence[str] | None = None,
        borders: bool = True,
    ) -> None:
        if caption:
            self.table_no += 1
            label = f"Table {self.chapter_no}.{self.table_no}: {caption}"
            self.paragraph(label, size=self.style.caption, bold=True, align="center", before=6, after=4, line=1.0, keep_next=True)
            self.entries.append(Entry("table", 1, label, label, "body"))
        size = size or self.style.table
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for column, width in zip(table.columns, widths, strict=True):
            column.width = Inches(width)
        if borders:
            _borders(table)
        for r, values in enumerate(rows):
            for c, value in enumerate(values):
                cell = table.cell(r, c)
                cell.width = Inches(widths[c])
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = ALIGN[(align[c] if align else "left")]
                is_head = header and r == 0
                self.runs(p, value, size=size, bold=is_head, color=WHITE if is_head else None)
                if is_head:
                    _shade(cell, "1F3A5F")
        if header:
            _repeat_header(table.rows[0])
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        self._fresh_page = False

    def contents(self, entries: Sequence[Entry], *, size: float | None = None, width: float = 6.27) -> None:
        size = size or self.style.body
        head = self.doc.add_paragraph()
        head.paragraph_format.tab_stops.add_tab_stop(Inches(width), WD_TAB_ALIGNMENT.RIGHT)
        self.runs(head, "\t**Page No.**", size=size)
        for entry in entries:
            p = self.doc.add_paragraph()
            fmt = p.paragraph_format
            fmt.space_after = Pt(3)
            fmt.line_spacing = 1.0
            fmt.tab_stops.add_tab_stop(Inches(width), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            fmt.left_indent = Inches(0.35 * (entry.level - 1))
            page = self.pages.get(entry.label, "0")
            text = f"**{entry.text}**\t{page}" if entry.kind == "chapter" else f"{entry.text}\t{page}"
            self.runs(p, text, size=size)
        self._fresh_page = False

    def set_properties(self, *, title: str, subject: str, author: str, keywords: str = "") -> None:
        props = self.doc.core_properties
        props.title = title
        props.subject = subject
        props.author = author
        props.keywords = keywords
        props.comments = ""
        props.last_modified_by = author

    def save(self, path: Path) -> None:
        self.doc.save(str(path))

    # ----- helpers -----------------------------------------------------------------

    def _extract_logo(self) -> Path | None:
        shapes = self.doc.inline_shapes
        if not len(shapes):
            return None
        blip = shapes[0]._inline.graphic.graphicData.pic.blipFill.blip
        part = self.doc.part.related_parts[blip.embed]
        out = Path(tempfile.gettempdir()) / f"niet_logo_{id(self)}.{part.partname.ext}"
        out.write_bytes(part.blob)
        return out

    def logo_picture(self, width: float) -> None:
        if self.logo is None:
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.add_run().add_picture(str(self.logo), width=Inches(width))


def set_font(run, name: str, size: float, bold: bool, italic: bool, color: RGBColor | None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)


def _shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _borders(table) -> None:
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "7F8FA6")
        borders.append(el)
    tblpr.append(borders)


def _repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


# ----- two-pass page numbering ---------------------------------------------------------


def to_pdf(docx: Path) -> Path:
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
        check=True,
        capture_output=True,
    )
    return docx.with_suffix(".pdf")


def page_texts(pdf: Path) -> list[str]:
    out = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], check=True, capture_output=True, text=True).stdout
    pages = out.split("\f")
    return pages[:-1] if pages and not pages[-1].strip() else pages


def printed_number(page: str) -> str | None:
    lines = [line.strip() for line in page.splitlines() if line.strip()]
    for line in reversed(lines[-3:]):
        if re.fullmatch(r"[ivxlc]+|\d+", line):
            return line
    return None


def locate(entries: Sequence[Entry], pdf: Path) -> dict[str, str]:
    """Map each entry label to the printed page number where it first appears."""
    pages = page_texts(pdf)
    norm = [{_norm(line) for line in page.splitlines() if line.strip()} for page in pages]
    found: dict[str, str] = {}
    start = 1
    for entry in entries:
        key = _norm(entry.label)
        for index in range(start, len(pages)):
            if key in norm[index] or any(
                len(line) >= 30 and key.startswith(line) for line in norm[index]
            ):
                number = printed_number(pages[index])
                if number:
                    found[entry.label] = number
                    start = index
                break
    return found


def _norm(text: str) -> str:
    return " ".join(text.replace("\t", " ").split()).upper()


def build_twice(make: Callable[[dict[str, str] | None], Builder], out: Path, *, rounds: int = 3) -> Path:
    pages: dict[str, str] | None = None
    for _ in range(rounds):
        builder = make(pages)
        builder.save(out)
        pdf = to_pdf(out)
        new_pages = locate(builder.entries, pdf)
        missing = [e.label for e in builder.entries if e.label not in new_pages]
        if new_pages == pages:
            if missing:
                print("not located:", missing)
            return pdf
        pages = new_pages
    return pdf
